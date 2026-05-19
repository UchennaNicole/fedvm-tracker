"""
poam_generator.py
-----------------
Auto-generates a Plan of Action & Milestones (POA&M) from scored vulnerability
findings. Output is formatted to align with OMB Memorandum M-02-01 and the
standard Federal POA&M template used in FISMA reporting.

The generated POA&M includes:
    - Weakness ID & description
    - Control mapping (NIST SP 800-53 Rev 5)
    - Point of Contact (system/app owner)
    - Resources required
    - Scheduled completion date (based on remediation tier SLA)
    - Milestone with completion dates
    - Status (Open / In Progress / Closed)
    - Risk rating (from remediation tier)
    - Identified weakness date
    - Delay reason (if overdue)

Output formats:
    - Excel (.xlsx) — primary Federal deliverable
    - CSV           — for import into GRC tools (Drata, Hyperproof, etc.)

Usage:
    from poam.poam_generator import POAMGenerator
    gen = POAMGenerator(system_name="DOT-OCIO-MAIN", system_owner="John Smith")
    gen.generate(scored_findings, output_path="reports/poam_output.xlsx")
"""

import csv
import logging
from datetime import date, datetime
from pathlib import Path
from typing import Optional

log = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# NIST 800-53 Rev 5 control mapping by vulnerability category
# ---------------------------------------------------------------------------

# Maps finding name keywords → primary NIST 800-53 controls
# Extend this map as your environment's vulnerability profile grows.
CONTROL_MAP = {
    "patch":        ["SI-2", "SI-2(2)"],
    "update":       ["SI-2", "SI-2(2)"],
    "unpatched":    ["SI-2", "SI-2(2)"],
    "smb":          ["SI-2", "CM-7", "SC-7"],
    "rdp":          ["AC-17", "SI-2", "CM-7"],
    "ssl":          ["SC-8", "SC-28", "SI-2"],
    "tls":          ["SC-8", "SC-28"],
    "cipher":       ["SC-8", "SC-17"],
    "certificate":  ["SC-17", "IA-5"],
    "telnet":       ["CM-7", "SC-8", "AC-17"],
    "snmp":         ["CM-7", "IA-3"],
    "log4":         ["SI-2", "SI-10"],
    "apache":       ["SI-2", "CM-7"],
    "tomcat":       ["SI-2", "CM-7"],
    "mysql":        ["SI-2", "AC-6"],
    "enumeration":  ["IA-5", "AC-7"],
    "disclosure":   ["AC-3", "AC-6", "SI-12"],
    "injection":    ["SI-10", "SI-2"],
    "spooler":      ["SI-2", "CM-7", "AC-6"],
    "default":      ["IA-5", "CM-6"],
}

FALLBACK_CONTROLS = ["RA-5", "SI-2"]

TIER_RISK_RATING = {
    "P1": "Very High",
    "P2": "High",
    "P3": "Moderate",
    "P4": "Low",
}

STATUS_OPEN        = "Open"
STATUS_IN_PROGRESS = "In Progress"
STATUS_CLOSED      = "Closed"


# ---------------------------------------------------------------------------
# Generator
# ---------------------------------------------------------------------------

class POAMGenerator:
    """
    Generates a Federal-standard POA&M from scored vulnerability findings.

    Args:
        system_name:     Name of the information system (for header/metadata).
        system_owner:    Name of the system owner (Federal POC).
        organization:    Organization name (e.g., "DOT OCIO").
        fiscal_year:     FY for POA&M tracking (defaults to current year).
    """

    def __init__(
        self,
        system_name:  str = "Information System",
        system_owner: str = "System Owner",
        organization: str = "Department of Transportation",
        fiscal_year:  Optional[int] = None,
    ):
        self.system_name  = system_name
        self.system_owner = system_owner
        self.organization = organization
        self.fiscal_year  = fiscal_year or datetime.now().year

    # ------------------------------------------------------------------
    # Public interface
    # ------------------------------------------------------------------

    def generate(
        self,
        findings: list[dict],
        output_path: str = "reports/poam_output.xlsx",
        also_csv: bool = True,
    ) -> str:
        """
        Generate POA&M from findings list.

        Args:
            findings:     Scored findings from PriorityScorer.
            output_path:  Output path for the Excel file.
            also_csv:     Also write a CSV version alongside the Excel file.

        Returns:
            Path to the generated Excel file.
        """
        poam_rows = [self._build_poam_row(f, idx + 1) for idx, f in enumerate(findings)]

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        # Write Excel
        self._write_excel(poam_rows, output_path)

        # Optionally write CSV
        if also_csv:
            csv_path = output_path.with_suffix(".csv")
            self._write_csv(poam_rows, csv_path)
            log.info(f"CSV POA&M written: {csv_path}")

        log.info(f"POA&M generated: {output_path} ({len(poam_rows)} entries)")
        return str(output_path)

    # ------------------------------------------------------------------
    # Row builder
    # ------------------------------------------------------------------

    def _build_poam_row(self, finding: dict, seq: int) -> dict:
        """Build a single POA&M row from a scored finding dict."""
        today         = date.today()
        tier          = finding.get("remediation_tier", "P3")
        sla_days      = finding.get("tier_sla_days", 90)
        sla_due       = finding.get("sla_due_date", "")
        severity      = finding.get("severity", "unknown").capitalize()
        risk_rating   = TIER_RISK_RATING.get(tier, "Moderate")
        controls      = self._map_controls(finding.get("name", ""))
        overdue       = self._is_overdue(sla_due, today)
        delay_reason  = "Awaiting system owner scheduling" if overdue else ""
        milestone     = self._build_milestone(tier, sla_due)

        weakness_id = (
            f"FY{self.fiscal_year}-"
            f"{finding.get('severity', 'UNK')[0].upper()}-"
            f"{seq:04d}"
        )

        return {
            "Weakness ID":                weakness_id,
            "System Name":                self.system_name,
            "Weakness Name":              finding.get("name", "Unknown Vulnerability"),
            "Weakness Description":       self._truncate(finding.get("synopsis", ""), 500),
            "CVE":                        finding.get("cve", "N/A"),
            "CVSS Score":                 finding.get("cvss_score", 0.0),
            "Risk Rating":                risk_rating,
            "Remediation Tier":           tier,
            "SLA (Days)":                 sla_days,
            "Affected Host(s)":           finding.get("host", ""),
            "Port / Protocol":            f"{finding.get('port', '')} / {finding.get('protocol', '')}",
            "NIST 800-53 Control(s)":     ", ".join(controls),
            "Point of Contact":           self.system_owner,
            "Organization":               self.organization,
            "Date Identified":            finding.get("scan_date", today.isoformat()),
            "Scheduled Completion Date":  sla_due,
            "Milestone":                  milestone,
            "Status":                     STATUS_OPEN,
            "Days Open":                  finding.get("days_open", 0),
            "Overdue":                    "YES" if overdue else "NO",
            "Delay Reason":               delay_reason,
            "KEV Match":                  "YES" if finding.get("kev_match") else "NO",
            "Exploit Available":          "YES" if finding.get("exploit_available") else "NO",
            "Asset Tier":                 finding.get("asset_tier", "unknown"),
            "Risk Priority Score":        finding.get("risk_priority_score", 0),
            "Recommended Solution":       self._truncate(finding.get("solution", ""), 500),
            "References":                 finding.get("see_also", ""),
        }

    # ------------------------------------------------------------------
    # Writers
    # ------------------------------------------------------------------

    def _write_excel(self, rows: list[dict], path: Path) -> None:
        """Write POA&M to Excel with header formatting."""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            log.warning("openpyxl not installed. Run: pip install openpyxl")
            log.warning("Falling back to CSV only.")
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"POA&M FY{self.fiscal_year}"

        if not rows:
            wb.save(path)
            return

        headers = list(rows[0].keys())

        # Header styling
        header_fill   = PatternFill("solid", fgColor="1F3864")
        header_font   = Font(bold=True, color="FFFFFF", size=10)
        header_align  = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border   = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"),
        )

        # Severity row fills
        SEVERITY_FILLS = {
            "Very High": PatternFill("solid", fgColor="C00000"),
            "High":      PatternFill("solid", fgColor="FF0000"),
            "Moderate":  PatternFill("solid", fgColor="FFC000"),
            "Low":       PatternFill("solid", fgColor="FFFF00"),
        }

        # Write headers
        for col_idx, header in enumerate(headers, start=1):
            cell             = ws.cell(row=1, column=col_idx, value=header)
            cell.font        = header_font
            cell.fill        = header_fill
            cell.alignment   = header_align
            cell.border      = thin_border

        ws.row_dimensions[1].height = 30

        # Write data rows
        for row_idx, row in enumerate(rows, start=2):
            risk_rating = row.get("Risk Rating", "")
            row_fill    = SEVERITY_FILLS.get(risk_rating)

            for col_idx, header in enumerate(headers, start=1):
                cell          = ws.cell(row=row_idx, column=col_idx, value=row[header])
                cell.border   = thin_border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_fill and header == "Risk Rating":
                    cell.fill = row_fill
                    cell.font = Font(bold=True, color="FFFFFF" if risk_rating in ("Very High", "High") else "000000")

        # Auto-width columns (capped at 60)
        for col in ws.columns:
            max_len = max(
                (len(str(cell.value or "")) for cell in col),
                default=10
            )
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 60)

        # Freeze header row
        ws.freeze_panes = "A2"

        # Add metadata sheet
        meta_ws = wb.create_sheet("Metadata")
        meta_ws.append(["Field", "Value"])
        meta_ws.append(["System Name",    self.system_name])
        meta_ws.append(["System Owner",   self.system_owner])
        meta_ws.append(["Organization",   self.organization])
        meta_ws.append(["Fiscal Year",    f"FY{self.fiscal_year}"])
        meta_ws.append(["Generated Date", date.today().isoformat()])
        meta_ws.append(["Total Findings", len(rows)])
        meta_ws.append(["Framework",      "NIST SP 800-53 Rev 5 / FISMA"])

        wb.save(path)

    def _write_csv(self, rows: list[dict], path: Path) -> None:
        if not rows:
            return
        with open(path, "w", newline="", encoding="utf-8") as fh:
            writer = csv.DictWriter(fh, fieldnames=rows[0].keys())
            writer.writeheader()
            writer.writerows(rows)

    # ------------------------------------------------------------------
    # Helpers
    # ------------------------------------------------------------------

    def _map_controls(self, vuln_name: str) -> list:
        """Map vulnerability name keywords to NIST 800-53 controls."""
        name_lower = vuln_name.lower()
        for keyword, controls in CONTROL_MAP.items():
            if keyword in name_lower:
                return controls
        return FALLBACK_CONTROLS

    @staticmethod
    def _is_overdue(sla_due: str, today: date) -> bool:
        if not sla_due:
            return False
        try:
            due = date.fromisoformat(sla_due)
            return today > due
        except ValueError:
            return False

    @staticmethod
    def _build_milestone(tier: str, due_date: str) -> str:
        milestones = {
            "P1": f"1. Validate scan findings — Day 3\n"
                  f"2. Coordinate patch deployment with sys admin — Day 7\n"
                  f"3. Apply patch / implement mitigation — Day 12\n"
                  f"4. Re-scan to confirm remediation — Day 15\n"
                  f"5. Close POA&M entry — {due_date}",
            "P2": f"1. Validate and triage findings — Day 5\n"
                  f"2. Assign to system owner — Day 7\n"
                  f"3. Schedule patch in change management — Day 14\n"
                  f"4. Apply patch — Day 25\n"
                  f"5. Re-scan and close — {due_date}",
            "P3": f"1. Validate findings — Week 1\n"
                  f"2. Assign to system owner — Week 2\n"
                  f"3. Remediation scheduled — Week 6\n"
                  f"4. Patch applied — Week 10\n"
                  f"5. Close POA&M — {due_date}",
            "P4": f"1. Document accepted risk or remediation path — Month 1\n"
                  f"2. Schedule remediation — Month 3\n"
                  f"3. Apply fix — Month 5\n"
                  f"4. Close POA&M — {due_date}",
        }
        return milestones.get(tier, f"Remediate by {due_date}")

    @staticmethod
    def _truncate(text: str, max_len: int) -> str:
        return text if len(text) <= max_len else text[:max_len - 3] + "..."


# ---------------------------------------------------------------------------
# CLI entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    import sys
    sys.path.insert(0, str(Path(__file__).parent.parent))

    from parsers.nessus_csv_parser import NessusParser
    from scoring.priority_scorer   import PriorityScorer

    filepath = sys.argv[1] if len(sys.argv) > 1 else "data/sample_nessus_export.csv"

    findings = NessusParser(filepath).parse()
    scored   = PriorityScorer(offline_mode=True).score_all(findings)

    gen = POAMGenerator(
        system_name  = "DOT-OCIO-PROD",
        system_owner = "Jane Smith",
        organization = "Department of Transportation OCIO",
    )
    out = gen.generate(scored, output_path="reports/poam_output.xlsx")
    print(f"\nPOA&M generated: {out}")
